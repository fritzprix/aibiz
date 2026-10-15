"""회의록 마크다운 → Word(docx). 사용: md2docx.py <입력.md> [출력.docx]

지원: 머리말(---) 건너뜀, # ~ ####, - / * 글머리(들여쓰기 2칸=한 단계), 1. 번호, - [ ] 체크,
      | 표 |, **굵게**, > 인용, --- 구분선.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

FONT = "맑은 고딕"


def set_font(style, size=None):
    style.font.name = FONT
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size:
        style.font.size = Pt(size)


def add_runs(par, text):
    for i, piece in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if piece:
            r = par.add_run(piece)
            r.bold = i % 2 == 1


def convert(md_path, out_path):
    lines = Path(md_path).read_text(encoding="utf-8").splitlines()
    if lines and lines[0].strip() == "---":
        end = next((i for i in range(1, len(lines)) if lines[i].strip() == "---"), 0)
        lines = lines[end + 1:]
    doc = Document()
    set_font(doc.styles["Normal"], 10.5)
    for s in ("Heading 1", "Heading 2", "Heading 3", "Heading 4", "Title"):
        set_font(doc.styles[s])
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()
        if not s or s.startswith("<!--"):
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                t.style = "Table Grid"
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, r in enumerate(rows):
                    for ci, c in enumerate(r):
                        p = t.cell(ri, ci).paragraphs[0]
                        add_runs(p, c.replace("**", "") if ri == 0 else c)
                        if ri == 0:
                            for run in p.runs:
                                run.bold = True
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            level = len(m.group(1))
            h = doc.add_heading(level=0 if level == 1 else level - 1)
            add_runs(h, m.group(2))
            i += 1
            continue
        if re.fullmatch(r"-{3,}", s):
            doc.add_paragraph("─" * 40).runs[0].font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
            i += 1
            continue
        indent = (len(line) - len(line.lstrip(" "))) // 2
        m = re.match(r"^[-*]\s+\[( |x|X)\]\s+(.*)", s)
        if m:
            p = doc.add_paragraph(style="List Bullet" if indent == 0 else "List Bullet 2")
            add_runs(p, ("☑ " if m.group(1).lower() == "x" else "☐ ") + m.group(2))
            i += 1
            continue
        m = re.match(r"^[-*]\s+(.*)", s)
        if m:
            style = ["List Bullet", "List Bullet 2", "List Bullet 3"][min(indent, 2)]
            add_runs(doc.add_paragraph(style=style), m.group(1))
            i += 1
            continue
        m = re.match(r"^\d+[.)]\s+(.*)", s)
        if m:
            add_runs(doc.add_paragraph(style="List Number"), m.group(1))
            i += 1
            continue
        if s.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            add_runs(p, s.lstrip("> "))
            i += 1
            continue
        add_runs(doc.add_paragraph(), s)
        i += 1
    doc.save(out_path)


if __name__ == "__main__":
    for _s in (sys.stdout, sys.stderr):
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".docx")
    convert(src, dst)
    print(f"DOCX: {dst}")
